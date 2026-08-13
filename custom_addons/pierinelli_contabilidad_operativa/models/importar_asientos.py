import base64
import io
from datetime import date, datetime

import openpyxl
import xlsxwriter

from odoo import fields, models, _
from odoo.exceptions import UserError, ValidationError


PLANTILLA_HEADERS = [
    'Referencia', 'Fecha', 'Código de diario', 'Código de cuenta', 'Glosa',
    'Debe', 'Haber', 'Tipo de operación', 'Contacto o RUC',
]
PLANTILLA_HEADERS_NORMALIZED = [
    'referencia', 'fecha', 'código de diario', 'código de cuenta', 'glosa',
    'debe', 'haber', 'tipo de operación', 'contacto o ruc',
]


class DescargarPlantillaAsientos(models.TransientModel):
    _name = 'pierinelli.descargar.plantilla.asientos'
    _description = 'Descargar plantilla de importación de asientos'

    archivo = fields.Binary('Plantilla Excel', readonly=True)
    archivo_nombre = fields.Char('Nombre de archivo', readonly=True)

    def default_get(self, fields_list):
        values = super().default_get(fields_list)
        output = io.BytesIO()
        workbook = xlsxwriter.Workbook(output, {'in_memory': True})
        main = workbook.add_worksheet('Asientos')
        instructions = workbook.add_worksheet('Instrucciones')
        header_format = workbook.add_format({
            'bold': True, 'font_color': '#FFFFFF', 'bg_color': '#1F4E78',
            'border': 1, 'align': 'center',
        })
        example_format = workbook.add_format({'font_color': '#7F7F7F'})
        date_format = workbook.add_format({'num_format': 'dd/mm/yyyy',
                                            'font_color': '#7F7F7F'})
        for column, header in enumerate(PLANTILLA_HEADERS):
            main.write(0, column, header, header_format)
        # Ejemplo claramente no importable: obliga a reemplazar los códigos de
        # diario y cuenta antes de usar la plantilla.
        example = [
            'REEMPLAZAR-001', date.today(), 'DIARIO_CODIGO', 'CUENTA_DEBE',
            'Glosa de ejemplo', 1000.00, 0.00, 'provision', 'RUC_O_NOMBRE_EXISTENTE',
        ]
        example_credit = [
            'REEMPLAZAR-001', date.today(), 'DIARIO_CODIGO', 'CUENTA_HABER',
            'Glosa de ejemplo', 0.00, 1000.00, 'provision', 'RUC_O_NOMBRE_EXISTENTE',
        ]
        for row_number, row in enumerate((example, example_credit), start=1):
            for column, value in enumerate(row):
                if column == 1:
                    main.write_datetime(row_number, column, value, date_format)
                else:
                    main.write(row_number, column, value, example_format)
        main.freeze_panes(1, 0)
        main.autofilter(0, 0, 0, len(PLANTILLA_HEADERS) - 1)
        for column, width in enumerate((22, 13, 18, 18, 36, 14, 14, 18, 30)):
            main.set_column(column, column, width)

        instructions.set_column(0, 0, 115)
        notes = [
            'PLANTILLA DE IMPORTACIÓN DE ASIENTOS',
            '1. Elimine las dos filas grises de ejemplo antes de importar.',
            '2. Cada fila es una línea contable. Las líneas del mismo asiento deben tener la misma Referencia, Fecha y Código de diario.',
            '3. Fecha: use dd/mm/aaaa. El Código de diario y Código de cuenta deben existir exactamente en Odoo.',
            '4. Debe y Haber son números; una línea solo debe tener importe en una de las dos columnas.',
            '5. Cada Referencia debe cuadrar: total Debe igual a total Haber. La importación crea únicamente borradores, nunca publica asientos.',
            '6. Tipo de operación es opcional: venta, compra, anticipo, detraccion, retencion, nota_credito, caja_chica, provision, depreciacion o ajuste.',
            '7. Contacto o RUC es opcional. Si se registra, debe coincidir exactamente con un contacto existente por RUC o nombre. Se asigna a las líneas y, si todo el asiento usa el mismo tercero, también a la cabecera.',
            '8. Primero pruebe con un asiento pequeño en una base de pruebas. Corrija cualquier error indicado antes de volver a subir el archivo.',
        ]
        title_format = workbook.add_format({'bold': True, 'font_size': 14,
                                             'font_color': '#1F4E78'})
        for index, note in enumerate(notes):
            instructions.write(index, 0, note,
                               title_format if index == 0 else None)
        workbook.close()
        values.update({
            'archivo': base64.b64encode(output.getvalue()),
            'archivo_nombre': 'plantilla_importacion_asientos.xlsx',
        })
        return values


class ImportarAsientos(models.TransientModel):
    _name = 'pierinelli.importar.asientos'
    _description = 'Importar asientos desde plantilla Excel'

    archivo = fields.Binary('Archivo Excel', required=True)
    archivo_nombre = fields.Char('Nombre de archivo', required=True)

    def _error(self, row, message):
        raise ValidationError(_('Fila %(row)s: %(message)s') % {
            'row': row, 'message': message})

    def _date_value(self, value, row):
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        if isinstance(value, str):
            for fmt in ('%d/%m/%Y', '%Y-%m-%d'):
                try:
                    return datetime.strptime(value.strip(), fmt).date()
                except ValueError:
                    continue
        self._error(row, _('fecha inválida; use dd/mm/aaaa.'))

    def _amount(self, value, label, row):
        if value in (None, ''):
            return 0.0
        try:
            amount = float(value)
        except (TypeError, ValueError):
            self._error(row, _('%s debe ser numérico.') % label)
        if amount < 0:
            self._error(row, _('%s no puede ser negativo.') % label)
        return round(amount, 2)

    def _partner_from_value(self, value, row):
        identifier = str(value or '').strip()
        if not identifier:
            return self.env['res.partner']
        partner = self.env['res.partner'].search([('vat', '=', identifier)], limit=1)
        if not partner:
            partner = self.env['res.partner'].search([('name', '=', identifier)], limit=1)
        if not partner:
            self._error(row, _(
                'no existe el contacto o RUC %(identifier)s. Créalo primero en Contactos.') % {
                    'identifier': identifier})
        return partner

    def _parse_file(self):
        self.ensure_one()
        try:
            content = base64.b64decode(self.archivo)
            book = openpyxl.load_workbook(io.BytesIO(content),
                                          read_only=True, data_only=True)
        except Exception as error:
            raise UserError(_(
                'No se pudo leer el archivo. Descarga la plantilla y súbela '
                'sin cambiar su estructura. Detalle: %s') % error)
        sheet = book['Asientos'] if 'Asientos' in book.sheetnames else book.active
        rows = sheet.iter_rows(values_only=True)
        try:
            headers = next(rows)
        except StopIteration:
            raise ValidationError(_('El archivo no tiene encabezados.'))
        normalized = [str(value or '').strip().lower() for value in headers]
        if normalized[:len(PLANTILLA_HEADERS_NORMALIZED)] != PLANTILLA_HEADERS_NORMALIZED:
            raise ValidationError(_(
                'Los encabezados no coinciden. Descarga nuevamente la plantilla '
                'y conserva: %s.') % ', '.join(PLANTILLA_HEADERS))

        lines = []
        for excel_row, values in enumerate(rows, start=2):
            if not any(value not in (None, '') for value in values):
                continue
            values = list(values) + [None] * (len(PLANTILLA_HEADERS) - len(values))
            reference = str(values[0] or '').strip()
            journal_code = str(values[2] or '').strip()
            account_code = str(values[3] or '').strip()
            description = str(values[4] or '').strip()
            if not reference or not journal_code or not account_code or not description:
                self._error(excel_row, _(
                    'Referencia, Código de diario, Código de cuenta y Glosa son obligatorios.'))
            debit = self._amount(values[5], _('Debe'), excel_row)
            credit = self._amount(values[6], _('Haber'), excel_row)
            if (debit and credit) or not (debit or credit):
                self._error(excel_row, _(
                    'ingresa importe solo en Debe o solo en Haber.'))
            operation_type = str(values[7] or '').strip()
            if operation_type and operation_type not in dict(
                    self.env['account.move']._fields['tipo_operacion_contable'].selection):
                self._error(excel_row, _(
                    'Tipo de operación no es válido. Déjalo vacío o usa uno de los valores de la plantilla.'))
            lines.append({
                'row': excel_row, 'reference': reference,
                'date': self._date_value(values[1], excel_row),
                'journal_code': journal_code, 'account_code': account_code,
                'name': description, 'debit': debit, 'credit': credit,
                'operation_type': operation_type or False,
                'partner_identifier': str(values[8] or '').strip(),
            })
        if not lines:
            raise ValidationError(_('No hay líneas para importar.'))
        return lines

    def action_importar(self):
        self.ensure_one()
        lines = self._parse_file()
        company = self.env.company
        journals = {}
        accounts = {}
        partners = {}
        grouped = {}
        for line in lines:
            journal = journals.get(line['journal_code'])
            if journal is None:
                journal = self.env['account.journal'].search([
                    ('code', '=', line['journal_code']),
                    ('company_id', '=', company.id)], limit=1)
                journals[line['journal_code']] = journal
            if not journal:
                self._error(line['row'], _(
                    'no existe el diario %(code)s para la compañía actual.') %
                    {'code': line['journal_code']})
            if journal.type != 'general':
                self._error(line['row'], _(
                    'el diario %(code)s debe ser de tipo Misceláneo/General para importar asientos.') %
                    {'code': line['journal_code']})
            account = accounts.get(line['account_code'])
            if account is None:
                account = self.env['account.account'].search([
                    ('code', '=', line['account_code']),
                    ('company_ids', 'in', company.id)], limit=1)
                accounts[line['account_code']] = account
            if not account:
                self._error(line['row'], _(
                    'no existe la cuenta %(code)s para la compañía actual.') %
                    {'code': line['account_code']})
            identifier = line['partner_identifier']
            if identifier not in partners:
                partners[identifier] = self._partner_from_value(identifier, line['row'])
            line['partner_id'] = partners[identifier].id
            key = (line['reference'], line['date'], journal.id,
                   line['operation_type'])
            grouped.setdefault(key, []).append((line, account))

        for (reference, _date, journal_id, operation_type), move_lines in grouped.items():
            debit = round(sum(line['debit'] for line, account in move_lines), 2)
            credit = round(sum(line['credit'] for line, account in move_lines), 2)
            if debit != credit:
                row = move_lines[0][0]['row']
                self._error(row, _(
                    'el asiento %(ref)s no cuadra: Debe %(debit).2f / Haber %(credit).2f.') % {
                        'ref': reference, 'debit': debit, 'credit': credit})

        moves = self.env['account.move']
        for (reference, move_date, journal_id, operation_type), move_lines in grouped.items():
            partner_ids = {line['partner_id'] for line, account in move_lines if line['partner_id']}
            moves |= self.env['account.move'].create({
                'move_type': 'entry', 'journal_id': journal_id,
                'date': move_date, 'ref': reference,
                'tipo_operacion_contable': operation_type,
                'partner_id': partner_ids.pop() if len(partner_ids) == 1 else False,
                'line_ids': [(0, 0, {
                    'account_id': account.id, 'name': line['name'],
                    'debit': line['debit'], 'credit': line['credit'],
                    'partner_id': line['partner_id'],
                }) for line, account in move_lines],
            })
        return {
            'type': 'ir.actions.act_window',
            'name': _('%s asiento(s) importado(s) en borrador') % len(moves),
            'res_model': 'account.move', 'view_mode': 'list,form',
            'domain': [('id', 'in', moves.ids)],
            'context': {'create': False},
        }
